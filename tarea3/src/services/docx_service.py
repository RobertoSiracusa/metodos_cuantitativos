"""
Servicio de Compilacion del Informe Tecnico Oficial en Word (.docx)
Capa de Servicios — Metodos Cuantitativos
Universidad Jose Antonio Paez
"""

from pathlib import Path
from typing import Dict, Optional, Union
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from src.utils.config import OUTPUTS_DIR


class DocxReportService:
    """
    Genera el Entregable 3 oficial: Informe Tecnico en formato .docx
    con membrete formal de la Universidad Jose Antonio Paez,
    desarrollo matematico, tabla de metricas, imagen de interfaz y diagnostico.
    """

    @staticmethod
    def _aplicar_sombreado(celda, color_hex="F1F5F9"):
        tc_pr = celda._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        tc_pr.append(shd)

    @classmethod
    def compilar_informe_docx(
        cls,
        metricas: Dict,
        diagnostico_ia: str,
        ruta_imagen: Optional[Union[str, Path]] = None,
        destino_docx: Union[str, Path] = "Informe_Tecnico_Simulador_Redes.docx",
    ) -> Path:
        """Compila y guarda el documento Word formal."""
        ruta_out = Path(destino_docx)
        if not ruta_out.is_absolute() and len(ruta_out.parts) == 1:
            ruta_out = OUTPUTS_DIR / ruta_out.name
        ruta_out.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()

        # Ajuste de margenes estandar (1 pulgada = 2.54 cm)
        for sec in doc.sections:
            sec.top_margin = Inches(1.0)
            sec.bottom_margin = Inches(1.0)
            sec.left_margin = Inches(1.0)
            sec.right_margin = Inches(1.0)

        # ----------------------------------------------------------------------
        # 1. PORTADA Y MEMBRETE INSTITUCIONAL
        # ----------------------------------------------------------------------
        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p_head.add_run("UNIVERSIDAD JOSE ANTONIO PAEZ\n")
        r1.font.name = "Arial"
        r1.font.size = Pt(14)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(15, 23, 42)

        r2 = p_head.add_run(
            "FACULTAD DE INGENIERIA — ESCUELA DE INGENIERIA EN COMPUTACION\n"
            "CATEDRA: METODOS CUANTITATIVOS Y SIMULACION\n\n"
        )
        r2.font.name = "Arial"
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(71, 85, 105)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_t = p_title.add_run("INFORME TECNICO:\nSIMULADOR DINAMICO DE REDES DE COMPUTADORAS\n")
        r_t.font.name = "Arial"
        r_t.font.size = Pt(16)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(37, 99, 235)

        r_sub = p_title.add_run(
            "Integracion de Teoria de Colas (M/M/1/K), Modelos de Inventario (s, Q)\n"
            "y Asignacion Optima con Algoritmo Hungaro\n\n"
        )
        r_sub.font.name = "Arial"
        r_sub.font.size = Pt(11)
        r_sub.font.italic = True
        r_sub.font.color.rgb = RGBColor(100, 116, 139)

        # ----------------------------------------------------------------------
        # 2. RESUMEN EJECUTIVO Y OBJETIVOS
        # ----------------------------------------------------------------------
        cls._agregar_titulo(doc, "1. Resumen Ejecutivo y Objetivos")
        doc.add_paragraph(
            "El presente proyecto describe el diseno, desarrollo y evaluacion cuantitativa de un simulador "
            "interactivo de redes de computadoras en tiempo real desarrollado en Python bajo el motor de eventos "
            "discretos SimPy y la biblioteca grafica Pygame. El software unifica tres modelos fundamentales de "
            "investigacion de operaciones:\n"
            "1. Teoria de Colas para la generacion estocastica de trafico Poisson y conmutacion en routers.\n"
            "2. Modelos de Inventario para la gestion de buffers de memoria y costos de ruptura por overflow.\n"
            "3. Modelo de Asignacion mediante Algoritmo Hungaro para optimizacion del enrutamiento dinamico."
        )

        # ----------------------------------------------------------------------
        # 3. MARCO MATEMATICO
        # ----------------------------------------------------------------------
        cls._agregar_titulo(doc, "2. Fundamentacion Matematica y Modelos Cuantitativos")

        cls._agregar_subtitulo(doc, "A. Teoria de Lineas de Espera (Modelo M/M/1/K)")
        doc.add_paragraph(
            "Las llegadas de paquetes siguen un Proceso de Poisson con tasa media lambda (arribos inter-llegada "
            "exponenciales Exp(lambda)). Los routers intermedios operan como canales de atencion exponencial con "
            "tasa mu. Se calculan en tiempo real:\n"
            "- L: Numero promedio de paquetes en el sistema.\n"
            "- Lq: Numero promedio de paquetes en cola de espera.\n"
            "- W: Tiempo medio de permanencia en la red.\n"
            "- Wq: Tiempo medio de espera en cola antes de ser atendido.\n"
            "- Cumplimiento de la Ley de Little: L = lambda_efectivo * W y Lq = lambda_efectivo * Wq."
        )

        cls._agregar_subtitulo(doc, "B. Modelo de Control de Inventario en Buffers")
        doc.add_paragraph(
            "Los buffers de memoria RAM se modelan bajo una politica de inventario finito (s, Q):\n"
            "- Capacidad maxima de almacenamiento S paquetes.\n"
            "- Umbral de reabastecimiento s: cuando el stock cae por debajo de s, se emite senal de control de flujo.\n"
            "- Costo de mantener en memoria (Holding Cost): Ch = $0.05 por paquete por segundo almacenado.\n"
            "- Costo de ruptura (Shortage Cost): Cs = $10.00 de penalizacion por paquete descartado por Buffer Overflow."
        )

        cls._agregar_subtitulo(doc, "C. Asignacion Optima con Algoritmo Hungaro")
        doc.add_paragraph(
            "En intervalos discretos Delta t = 1.0 s, se evalua la matriz de costos C_{ij} entre N flujos y M enlaces:\n"
            "C_{ij} = Latencia_Actual_{ij} + alpha * (Saturacion_Buffer_Nodo_j)\n"
            "donde alpha = 40.0 ms pondera la ocupacion del router de destino. Los enlaces caidos se penalizan "
            "con un costo prohibitivo de 10^6, forzando la reasignacion por rutas alternas."
        )

        # ----------------------------------------------------------------------
        # 4. TABLA DE RESULTADOS DE LA SIMULACION
        # ----------------------------------------------------------------------
        cls._agregar_titulo(doc, "3. Metricas Obtenidas en la Corrida de Evaluacion")

        tabla = doc.add_table(rows=1, cols=2)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = tabla.rows[0].cells
        hdr_cells[0].text = "Metrica Cuantitativa Evaluada"
        hdr_cells[1].text = "Valor Obtenido"
        cls._aplicar_sombreado(hdr_cells[0], "1E293B")
        cls._aplicar_sombreado(hdr_cells[1], "1E293B")
        for cell in hdr_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

        filas_datos = [
            ("Tiempo Total de Simulacion", f"{metricas.get('tiempo_simulacion_s', 0.0)} s"),
            ("Tasa de Llegada (lambda)", f"{metricas.get('lambda', 0.0):.1f} paq/s"),
            ("Tasa de Servicio por Servidor (mu)", f"{metricas.get('mu', 0.0):.1f} paq/s"),
            ("Capacidad de Buffer (S)", f"{metricas.get('capacidad_buffer_s', 50)} paquetes"),
            ("Umbral de Reabastecimiento (s)", f"{metricas.get('umbral_reabastecimiento_s', 10)} paquetes"),
            ("Total Paquetes Procesados", str(metricas.get("paquetes_procesados", 0))),
            ("Total Paquetes Perdidos (Overflow)", str(metricas.get("paquetes_perdidos", 0))),
            ("Tasa de Perdida de Paquetes", f"{metricas.get('tasa_perdida_pct', 0.0):.2f}%"),
            ("Tiempo Medio en Cola (Wq)", f"{metricas.get('tiempo_medio_cola_wq_s', 0.0):.4f} s"),
            ("Promedio de Paquetes en el Sistema (L)", f"{metricas.get('promedio_paquetes_sistema_l', 0.0):.2f}"),
            ("Promedio de Paquetes en Cola (Lq)", f"{metricas.get('promedio_paquetes_cola_lq', 0.0):.2f}"),
            ("Costo Total de Almacenamiento (RAM)", f"${metricas.get('costo_almacenamiento_usd', 0.0):.2f}"),
            ("Costo Total de Penalizacion (Ruptura)", f"${metricas.get('costo_penalizacion_usd', 0.0):.2f}"),
            ("Costo Global del Sistema", f"${metricas.get('costo_global_usd', 0.0):.2f}"),
        ]

        for i, (nombre, val) in enumerate(filas_datos):
            row = tabla.add_row()
            c0, c1 = row.cells
            c0.text = nombre
            c1.text = val
            if i % 2 == 1:
                cls._aplicar_sombreado(c0, "F8FAFC")
                cls._aplicar_sombreado(c1, "F8FAFC")

        doc.add_paragraph()

        # ----------------------------------------------------------------------
        # 5. CAPTURA DE INTERFAZ GRAFICA
        # ----------------------------------------------------------------------
        if ruta_imagen and Path(ruta_imagen).exists():
            cls._agregar_titulo(doc, "4. Evidencia Visual del Simulador Interactivo")
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(str(ruta_imagen), width=Inches(6.2))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cap = p_cap.add_run("Figura 1: Topologia de red dinamica con dashboard HUD en Pygame.")
            r_cap.font.italic = True
            r_cap.font.size = Pt(9)
            r_cap.font.color.rgb = RGBColor(100, 116, 139)

        # ----------------------------------------------------------------------
        # 6. DIAGNOSTICO DE INTELIGENCIA ARTIFICIAL Y CONCLUSIONES
        # ----------------------------------------------------------------------
        cls._agregar_titulo(doc, "5. Diagnostico Automatizado de Inteligencia Artificial")
        for bloque in diagnostico_ia.split("\n\n"):
            if bloque.strip():
                doc.add_paragraph(bloque.strip())

        doc.save(str(ruta_out))
        return ruta_out

    @staticmethod
    def _agregar_titulo(doc, texto: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(texto)
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)

    @staticmethod
    def _agregar_subtitulo(doc, texto: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(texto)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(37, 99, 235)
